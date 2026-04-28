from Interface.IService import IService
from docx import Document

class DocxService(IService):

    def __init__(self):
        self.files_directory = "Download_Files/"

    def read_file(self, file_name):
        path = self.files_directory + file_name
        print(path)
        document = Document(path)
        text = ""
        for paragraph in document.paragraphs:
            if paragraph.text:
                text += paragraph.text + "\n"
        text += self.__read_table(document)
        return text

    def __read_table(self, document):
        text = ""
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        return text