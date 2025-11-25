from Interface.IService import IService
from openpyxl import load_workbook


class ExcelService(IService):

    def __init__(self):
        self.files_directory = "Download_Files/"

    def read_file(self, file_name):

        path = self.files_directory + file_name

        workbook = load_workbook(path)
        sheet = workbook.active 

        all_text = ""
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
            all_text += row_text + "\n"
        
        file = open("Report.txt", "a", encoding="utf-8")
        file.write(all_text + "\n")
        file.close()
        return all_text
        
